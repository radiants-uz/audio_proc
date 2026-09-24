"""§5.9: отчёт .docx (python-docx) → .pdf (LibreOffice headless).
Шаблона от заказчика нет, поэтому документ собирается кодом; секции — по
плану: титул, резюме, темы, решения, задачи, участники, статистика,
приложение с транскриптом, колонтитул с дисклеймером."""

import io
import logging
import shutil
import subprocess
import tempfile
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.core.config import settings
from app.services.audio_processing.analysis.analyze import segment_language_and_text
from app.services.audio_processing.analysis.stats import compute_stats
from app.services.audio_processing.report.charts import languages_chart, speakers_chart
from app.services.audio_processing.report.i18n import strings

logger = logging.getLogger(__name__)

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_CT = "application/pdf"


def _hms(sec: float) -> str:
    sec = int(sec)
    return f"{sec // 3600:02d}:{(sec % 3600) // 60:02d}:{sec % 60:02d}"


def _table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return table


def build_docx(transcript: dict, analysis: dict, title: str | None, speaker_names: dict[str, str] | None = None) -> bytes:
    lang = analysis.get("report_lang", "ru")
    s = strings(lang)
    names = {sp["id"]: (sp.get("name") or sp["id"]) for sp in transcript.get("speakers", [])}
    names.update({k: v for k, v in (speaker_names or {}).items() if v})
    stats = analysis.get("stats", {})

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(10)

    # Титул
    heading = title or analysis.get("title") or transcript.get("title") or s["report"]
    h = doc.add_heading(heading, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    audio = transcript.get("audio", {})
    # все языки файла по времени речи (спикер может говорить на двух), а не только «главные»
    by_language = stats.get("by_language") or compute_stats(transcript).get("by_language", [])
    langs = ", ".join(x["language"] for x in by_language if x.get("language"))
    meta_lines = [
        f"{s['file']}: {audio.get('filename', '')}",
        f"{s['date']}: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}",
        f"{s['duration']}: {_hms(audio.get('duration_sec', 0))}",
        f"{s['participants']}: {len(transcript.get('speakers', []))}",
        f"{s['languages']}: {langs or s['none']}",
    ]
    for line in meta_lines:
        doc.add_paragraph(line)

    # Резюме
    doc.add_heading(s["summary"], level=1)
    doc.add_paragraph(analysis.get("summary") or s["none"])

    # Темы
    doc.add_heading(s["topics"], level=1)
    topics = analysis.get("topics") or []
    if not topics:
        doc.add_paragraph(s["none"])
    for t in topics:
        p = doc.add_paragraph()
        p.add_run(t.get("name", "")).bold = True
        doc.add_paragraph(t.get("summary", ""))

    # Решения
    doc.add_heading(s["decisions"], level=1)
    seg_start = {seg["id"]: seg["start"] for seg in transcript.get("segments", [])}

    def _time(ids):
        ids = [i for i in (ids or []) if i in seg_start]
        return _hms(seg_start[min(ids)]) if ids else s["none"]

    decisions = analysis.get("decisions") or []
    if decisions:
        _table(doc, [s["col_decision"], s["col_who"], s["col_time"]],
               [[d.get("text", ""), names.get(d.get("speaker") or "", d.get("speaker") or s["none"]), _time(d.get("segment_ids"))] for d in decisions],
               [4.3, 1.4, 0.8])
    else:
        doc.add_paragraph(s["none"])

    # Задачи
    doc.add_heading(s["action_items"], level=1)
    items = analysis.get("action_items") or []
    if items:
        _table(doc, [s["col_task"], s["col_assignee"], s["col_due"], s["col_time"]],
               [[a.get("text", ""), names.get(a.get("assignee") or "", a.get("assignee") or s["none"]), a.get("due") or s["none"], _time(a.get("segment_ids"))] for a in items],
               [3.5, 1.3, 0.9, 0.8])
    else:
        doc.add_paragraph(s["none"])

    # Участники
    doc.add_heading(s["speakers"], level=1)
    positions = {sp["id"]: sp for sp in analysis.get("speakers") or []}
    rows = []
    for st in stats.get("by_speaker", []):
        a = positions.get(st["id"], {})
        pos = a.get("position") or ""
        if a.get("role_hint"):
            pos = f"{a['role_hint']}. {pos}"
        rows.append([names.get(st["id"], st["id"]), st.get("language", ""), _hms(st["speech_sec"]),
                     f"{st['share'] * 100:.0f}%", pos])
    if rows:
        _table(doc, [s["col_speaker"], s["col_language"], s["col_speech"], s["col_share"], s["col_position"]], rows,
               [1.3, 0.6, 0.9, 0.6, 3.1])

    # Открытые вопросы
    oq = analysis.get("open_questions") or []
    if oq:
        doc.add_heading(s["open_questions"], level=1)
        for q in oq:
            doc.add_paragraph(q, style="List Bullet")

    # Статистика
    doc.add_heading(s["stats"], level=1)
    if stats.get("by_speaker"):
        doc.add_picture(io.BytesIO(speakers_chart(stats, s["chart_speakers"])), width=Inches(5.5))
    if stats.get("by_language"):
        doc.add_picture(io.BytesIO(languages_chart(stats, s["chart_languages"])), width=Inches(5.5))
    doc.add_paragraph(f"{s['overlaps_count']}: {stats.get('overlap_segments', 0)}")
    doc.add_paragraph(f"{s['segments']}: {stats.get('segments_count', 0)}")

    # Приложение — транскрипт
    if settings.report.include_full_transcript:
        doc.add_page_break()
        doc.add_heading(s["transcript"], level=1)
        rows = []
        for seg in transcript.get("segments", []):
            lang, text = segment_language_and_text(seg)
            if seg.get("overlap"):
                text = f"[{s['overlap']}] {text}"
            rows.append([_hms(seg["start"]), names.get(seg["speaker"], seg["speaker"]), lang, text])
        if rows:
            _table(doc, [s["col_time"], s["col_speaker"], s["col_language"], s["col_text"]], rows, [0.8, 1.2, 0.5, 4.0])

    # Колонтитул
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"{s['disclaimer']}  job {transcript.get('job_id', '')}  pipeline {transcript.get('meta', {}).get('pipeline_version', '')}"
    footer.runs[0].font.size = Pt(7)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_to_pdf(docx_bytes: bytes) -> bytes | None:
    """LibreOffice headless. None — если конвертация не удалась (warning pdf_failed)."""
    soffice = shutil.which(settings.report.libreoffice_bin) or shutil.which("libreoffice")
    if not soffice:
        logger.warning("LibreOffice not found (%s)", settings.report.libreoffice_bin)
        return None
    with tempfile.TemporaryDirectory(prefix="report-") as tmp:
        src = Path(tmp) / "report.docx"
        src.write_bytes(docx_bytes)
        try:
            subprocess.run(
                [soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/lo-profile",
                 "--convert-to", "pdf", "--outdir", tmp, str(src)],
                capture_output=True, text=True, timeout=300, check=True,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            logger.warning("docx → pdf failed: %s", getattr(e, "stderr", "") or e)
            return None
        pdf = Path(tmp) / "report.pdf"
        return pdf.read_bytes() if pdf.exists() else None
