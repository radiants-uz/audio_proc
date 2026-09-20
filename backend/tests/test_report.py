import io

from docx import Document

from app.services.audio_processing.report.build import build_docx
from tests.test_analysis import GOOD, TRANSCRIPT


def test_build_docx_all_languages():
    for lang in ("ru", "uz", "en"):
        analysis = {**GOOD, "report_lang": lang, "stats": {
            "total_speech_sec": 26, "speakers_count": 2, "segments_count": 4, "overlap_segments": 1, "avg_segment_sec": 6.5,
            "by_speaker": [{"id": "SPEAKER_00", "name": None, "language": "ru", "speech_sec": 18, "share": 0.69, "segments": 3},
                           {"id": "SPEAKER_01", "name": None, "language": "uz", "speech_sec": 8, "share": 0.31, "segments": 1}],
            "by_language": [{"language": "ru", "speech_sec": 18, "share": 0.69}, {"language": "uz", "speech_sec": 8, "share": 0.31}],
        }, "meta": {"llm": "x", "prompt_version": "1"}}
        data = build_docx(TRANSCRIPT, analysis, None, {"SPEAKER_01": "Bobur"})
        assert len(data) > 10_000
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Планёрка" in text
        cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        assert "Bobur" in cells and "Boshlaymiz." in " ".join(cells)
